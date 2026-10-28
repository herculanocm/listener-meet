"""A reunião em .docx (F5.1 item 7, U-011, A-18; D-012: só por ação do usuário, não fere a CP9): as mesmas seções do .md (Q4)
na mesma ordem: o título, a ficha (data e hora, duração, participantes), a ata, o resumo executivo e o pré-resumo (o
Markdown deles lido pelo parser puro de `rag/markdown.py`: títulos, listas, negrito e itálico), as ações, os marcadores e
a transcrição "[mm:ss] Nome — Cargo: texto" com o falante em negrito. Texto do usuário (falas, notas, ações) entra como
texto, sem virar Markdown.

Todo texto que entra no python-docx passa por `_xml_text` (R9-01): um caractere de controle que o XML proíbe, num título,
nome, dono ou na ata, faria o documento inteiro falhar; o título do arquivo (propriedade do .docx) é cortado em 255
caracteres, o limite do python-docx (o título da reunião não tem limite).

O `python-docx` (extra `docx`) só é importado dentro de `meeting_docx`, que roda numa `Task` da reunião aberta: o app abre
sem o extra, e sem ele o item do menu fica desabilitado com a dica (`docx_available`, que procura sem importar). O
documento sai em bytes; a escrita é a atômica do .md (`export.write_atomic`).

F5.3 (A-07, pergunta 8): os títulos das seções no idioma da ata (os textos do .md, `export.TEXTS`) e o idioma padrão do
texto, que guia a revisão ortográfica do Word, em `pt-BR` ou `en-US` (`config.language.docx_code`)."""
from __future__ import annotations

import importlib.util
import re
from datetime import UTC, datetime
from io import BytesIO
from typing import TYPE_CHECKING

from ..config.language import docx_code
from ..llm.prompts.common import stamp
from . import export
from .export import ExportMeeting
from .markdown import parse_markdown

if TYPE_CHECKING:
    from docx.document import Document
    from docx.text.paragraph import Paragraph

AUTHOR = "listener-meet"         # o modelo padrão do python-docx assina "python-docx"
CORE_PROPERTY_MAX = 255          # o python-docx recusa propriedade do arquivo maior que isto
BULLET_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")
NUMBER_INDENT_PT = 18            # recuo por nível da lista numerada (o número vai escrito: o Word continuaria a contagem)
SECTION_LEVEL = 1                # "Ata", "Ações", "Transcrição"…; os títulos de dentro da ata começam no 2
XML_FORBIDDEN = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")   # o que o XML 1.0 (e o lxml) recusa


def docx_available() -> bool:
    """O extra `docx` instalado, sem importá-lo (o menu decide se o item fica habilitado)."""
    return importlib.util.find_spec("docx") is not None


def _xml_text(text: str) -> str:
    """O texto como o XML do .docx aceita: sem os caracteres de controle de \\x00 a \\x1f (salvo a tabulação e as quebras
    \\n e \\r), sem surrogates soltos e sem U+FFFE e U+FFFF. O resto (acentos, emoji, "<" e "&") passa como está."""
    return XML_FORBIDDEN.sub("", text)


def _run(p: Paragraph, text: str, *, bold: bool = False, italic: bool = False) -> None:
    """Um trecho do parágrafo com o texto limpo; sem marca, vale a do estilo (os títulos já são negrito)."""
    r = p.add_run(_xml_text(text))
    if bold:
        r.bold = True
    if italic:
        r.italic = True


def _heading(doc: Document, text: str, level: int) -> Paragraph:
    p = doc.add_heading("", level=level)
    if text:
        _run(p, text)
    return p


def meeting_docx(m: ExportMeeting) -> bytes:
    """O .docx inteiro da reunião, em bytes."""
    import docx  # o extra `docx` só nesta ação (A-18): nunca no nível do módulo

    doc = docx.Document()
    _properties(doc, m)
    _heading(doc, m.shown_title, 0)
    for label, value in export.header_facts(m):
        p = doc.add_paragraph()
        _run(p, f"{label}: ", bold=True)
        _run(p, value)
    t = m.texts
    for title, md in ((t.minutes, m.minutes_md), (t.executive, m.executive), (t.pre_summary, m.pre_summary)):
        if md.strip():
            _heading(doc, title, SECTION_LEVEL)
            _markdown(doc, md)
    _list_section(doc, t.actions, [export.action_text(a, m.language) for a in m.actions])
    _list_section(doc, t.markers, [export.marker_text(time, note, m.language) for time, note in m.markers])
    if m.transcript:
        _heading(doc, t.transcript, SECTION_LEVEL)
        for line in m.transcript:
            p = doc.add_paragraph()
            _run(p, f"[{stamp(line.t_ms)}] ")
            _run(p, line.speaker, bold=True)
            _run(p, f": {line.text}")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _properties(doc: Document, m: ExportMeeting) -> None:
    """Título, autor e datas do arquivo (o modelo padrão traz "python-docx" e 2013) e o idioma padrão do texto, o da ata
    (a revisão ortográfica do Word; o modelo padrão vem em inglês dos EUA)."""
    from docx.oxml.ns import qn

    props = doc.core_properties
    now = datetime.now(UTC)
    props.title = _xml_text(m.shown_title)[:CORE_PROPERTY_MAX]
    props.author = props.last_modified_by = AUTHOR
    props.created = props.modified = now
    lang = doc.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:lang')}")
    if lang is not None:
        lang.set(qn("w:val"), docx_code(m.language))


def _list_section(doc: Document, title: str, items: list[str]) -> None:
    if items:
        _heading(doc, title, SECTION_LEVEL)
        for text in items:
            _run(doc.add_paragraph(style=BULLET_STYLES[0]), text)


def _markdown(doc: Document, md: str) -> None:
    """O Markdown da ata ou de um resumo: os títulos dele começam no nível 2 (debaixo da seção), qualquer que seja o primeiro
    nível usado ("##" na ata da LLM; "#" numa ata editada)."""
    from docx.shared import Pt

    blocks = parse_markdown(md)
    top = min((b.level for b in blocks if b.kind == "heading"), default=1)
    for block in blocks:
        p: Paragraph
        if block.kind == "heading":
            p = _heading(doc, "", max(SECTION_LEVEL + 1, min(9, block.level - top + SECTION_LEVEL + 1)))
        elif block.kind == "bullet":
            p = doc.add_paragraph(style=BULLET_STYLES[min(block.level, len(BULLET_STYLES) - 1)])
        else:
            p = doc.add_paragraph()
            if block.kind == "number":
                p.paragraph_format.left_indent = Pt(NUMBER_INDENT_PT * (block.level + 1))
                _run(p, f"{block.number}. ")
        for run in block.runs:
            _run(p, run.text, bold=run.bold, italic=run.italic)


__all__ = ["docx_available", "meeting_docx"]
